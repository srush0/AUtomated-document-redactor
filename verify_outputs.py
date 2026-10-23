import re
from pathlib import Path

import pymupdf
from docx import Document

OUTPUT = Path("output")

PATTERNS = {
    "email": r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
    "Indian phone": r"(?<!\w)(?:\+91[- ]?)?[6-9]\d{9}(?!\w)",
    "API key": r"\b(?:sk-[A-Za-z0-9_-]{12,}|AKIA[A-Z0-9]{16})\b",
    "employee ID": r"\bEMP[-_]?\d{3,}\b",
}


def extract_docx(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)

    for section in doc.sections:
        for container in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            parts.extend(p.text for p in container.paragraphs)
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        parts.extend(p.text for p in cell.paragraphs)

    return "\n".join(parts)


def extract_text(path):
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        with pymupdf.open(path) as pdf:
            return "\n".join(page.get_text() for page in pdf)

    if suffix == ".docx":
        return extract_docx(path)

    if suffix == ".txt":
        return path.read_text(encoding="utf-8")

    return None


def main():
    files = sorted(
        p for p in OUTPUT.glob("*_redacted.*")
        if p.suffix.lower() in {".pdf", ".docx", ".txt"}
    )

    if not files:
        print("FAIL: No redacted PDF, DOCX, or TXT files found.")
        return 1

    failures = 0
    warnings = 0

    for path in files:
        print(f"\n--- {path.name} ---")

        try:
            text = extract_text(path)
        except Exception as exc:
            print(f"FAIL: Could not inspect file ({type(exc).__name__}).")
            failures += 1
            continue

        if text is None:
            continue

        if not text.strip():
            print("WARN: No extractable text. Scanned/image content may be unchecked.")
            warnings += 1
            continue

        file_failed = False
        for label, pattern in PATTERNS.items():
            count = len(re.findall(pattern, text, flags=re.IGNORECASE))
            if count:
                print(f"FAIL: {label} pattern remains ({count} match(es)).")
                file_failed = True

        if file_failed:
            failures += 1
        else:
            print("PASS: No configured regex patterns found in extracted text.")

    print("\n--- OVERALL RESULT ---")
    if failures:
        print(f"FAIL: {failures} file(s) need review.")
        print(f"Warnings: {warnings}")
        return 1

    print("PASS: No configured regex patterns detected in readable output.")
    print(f"Warnings: {warnings}")
    print("LIMITATION: This does not certify complete redaction.")
    print("Names, images, comments, metadata, embedded files and other")
    print("unsupported document components may still contain sensitive data.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
