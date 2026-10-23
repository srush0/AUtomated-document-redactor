from pathlib import Path
import tempfile
import re
import sqlite3
import time
from datetime import datetime

import streamlit as st
import pymupdf
from docx import Document
from ai_redactor import PATTERNS

from document_redactor import (
    redact_pdf,
    redact_docx,
    redact_txt,
    find_sensitive_items,
)

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "redaction_history.db"

MIME_TYPES = {
    ".pdf": "application/pdf",
    ".docx": (
        "application/vnd.openxmlformats-officedocument."
        "wordprocessingml.document"
    ),
    ".txt": "text/plain",
}


# -------------------- HISTORY DATABASE --------------------

def init_history_db():
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS redaction_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                timestamp TEXT NOT NULL,
                filename TEXT NOT NULL,
                file_type TEXT NOT NULL,
                file_size_bytes INTEGER NOT NULL,
                candidates INTEGER DEFAULT 0,
                operations INTEGER DEFAULT 0,
                categories TEXT DEFAULT '',
                duration_seconds REAL DEFAULT 0,
                processing_status TEXT NOT NULL,
                verification_status TEXT NOT NULL
            )
        """)


def save_history(
    filename,
    file_type,
    file_size,
    candidates,
    operations,
    categories,
    duration,
    processing_status,
    verification_status,
):
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute("""
            INSERT INTO redaction_history (
                timestamp, filename, file_type, file_size_bytes,
                candidates, operations, categories, duration_seconds,
                processing_status, verification_status
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            datetime.now().astimezone().isoformat(timespec="seconds"),
            filename,
            file_type,
            file_size,
            candidates,
            operations,
            ", ".join(categories),
            round(duration, 3),
            processing_status,
            verification_status,
        ))


def load_history():
    with sqlite3.connect(DB_PATH) as conn:
        return conn.execute("""
            SELECT timestamp, filename, file_type, file_size_bytes,
                   candidates, operations, categories, duration_seconds,
                   processing_status, verification_status
            FROM redaction_history
            ORDER BY id DESC
            LIMIT 200
        """).fetchall()


init_history_db()


# -------------------- DOCUMENT HELPERS --------------------

def read_docx_text(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)

    for section in doc.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)

    return "\n".join(parts)


def extract_text(path, suffix):
    if suffix == ".pdf":
        with pymupdf.open(path) as pdf:
            return "\n".join(page.get_text() for page in pdf)
    if suffix == ".docx":
        return read_docx_text(path)
    if suffix == ".txt":
        return Path(path).read_text(encoding="utf-8")
    raise ValueError("Unsupported document type.")


def normalize(text):
    return re.sub(r"\s+", " ", text).casefold()


def detect_residual_patterns(text):
    found = {}
    for label, pattern in PATTERNS.items():
        count = len(re.findall(pattern, text, flags=re.IGNORECASE))
        if count:
            found[label] = count
    return found


# -------------------- STREAMLIT INTERFACE --------------------

st.set_page_config(
    page_title="Automated Document Redactor",
    page_icon="🔐",
    layout="wide",
)

st.title("🔐 Automated Document Redactor")
st.write(
    "Detect and redact sensitive information using local regex "
    "rules and Gemma through Ollama."
)

st.info(
    "Use dummy files while testing. Verification checks extracted text "
    "and configured patterns only. It does not certify the entire "
    "document as safe."
)

st.subheader("How redaction works")

steps = [
    "1. Extract readable text from the uploaded document.",
    "2. Detect sensitive information using regex rules and local Gemma.",
    "3. Redact detected information in the supported document format.",
    "4. Verify the output text and check for remaining configured patterns.",
]
for step in steps:
    st.write(step)

st.divider()
st.subheader("Redact a document")

uploaded = st.file_uploader(
    "Upload a document",
    type=["pdf", "docx", "txt"],
)

if uploaded:
    st.write(f"Selected file: **{uploaded.name}**")

    if st.button("Redact and Verify", type="primary"):
        started = time.perf_counter()

        suffix = Path(uploaded.name).suffix.lower()
        file_size = len(uploaded.getvalue())
        candidates = []
        detected = 0
        operations = 0
        categories = []
        verification_status = "NOT COMPLETED"
        processing_status = "FAILED"
        download_data = None
        remaining_count = 0
        residual_patterns = {}

        try:
            if suffix not in MIME_TYPES:
                raise ValueError("Unsupported file type.")

            with tempfile.TemporaryDirectory() as temp_dir:
                temp = Path(temp_dir)
                source = temp / f"input{suffix}"
                destination = temp / f"redacted{suffix}"
                source.write_bytes(uploaded.getvalue())

                with st.spinner(
                    "Extracting text and detecting sensitive information..."
                ):
                    original_text = extract_text(source, suffix)

                    if not original_text.strip():
                        raise ValueError(
                            "No extractable text found. Scanned PDFs "
                            "require OCR, which is not implemented yet."
                        )

                    # Count configured categories in the original text.
                    original_pattern_counts = detect_residual_patterns(
                        original_text
                    )
                    categories = list(original_pattern_counts.keys())

                    candidates = find_sensitive_items(original_text)

                    if suffix == ".pdf":
                        operations, detected = redact_pdf(
                            source, destination
                        )
                    elif suffix == ".docx":
                        operations, detected = redact_docx(
                            source, destination
                        )
                    else:
                        operations, detected = redact_txt(
                            source, destination
                        )

                if not destination.is_file():
                    raise RuntimeError(
                        "Redacted output was not created."
                    )

                with st.spinner("Verifying the generated document..."):
                    output_text = extract_text(destination, suffix)
                    normalized_output = normalize(output_text)

                    remaining = [
                        item for item in candidates
                        if normalize(item) in normalized_output
                    ]
                    remaining_count = len(remaining)
                    residual_patterns = detect_residual_patterns(
                        output_text
                    )

                if remaining or residual_patterns:
                    verification_status = "FAILED"
                    processing_status = "COMPLETED"
                else:
                    verification_status = "PASSED"
                    processing_status = "COMPLETED"
                    download_data = destination.read_bytes()

            duration = time.perf_counter() - started

            save_history(
                uploaded.name,
                suffix.upper().lstrip("."),
                file_size,
                len(candidates),
                operations,
                categories,
                duration,
                processing_status,
                verification_status,
            )

            st.divider()
            st.subheader("Redaction Results")

            col1, col2, col3, col4 = st.columns(4)
            col1.metric("Candidates", len(candidates))
            col2.metric("Redaction operations", operations)
            col3.metric("Processing time", f"{duration:.2f} sec")
            col4.metric("File size", f"{file_size / 1024:.1f} KB")

            st.write("**Processing steps completed**")
            st.write("✓ Text extraction")
            st.write("✓ Sensitive information detection")
            st.write("✓ Document redaction")
            st.write(
                "✓ Output verification"
                if verification_status == "PASSED"
                else "⚠ Verification found remaining patterns"
            )

            st.write("**Sensitive-data categories found in input**")
            if categories:
                for category in categories:
                    st.write(
                        f"• {category}: "
                        f"{original_pattern_counts[category]} match(es)"
                    )
            else:
                st.write(
                    "No configured regex categories were detected "
                    "in the original text."
                )

            st.write(f"**Verification status:** {verification_status}")

            if verification_status == "FAILED":
                st.error(
                    "VERIFICATION FAILED. Download is blocked because "
                    "detected text or configured patterns remain."
                )
                st.write("Remaining candidate count:", remaining_count)
                st.write("Residual patterns:", residual_patterns)
                st.warning(
                    "Review the detection and redaction logic. "
                    "Do not treat this output as safe."
                )
            else:
                st.success(
                    "TEXT CHECK PASSED: no detected candidate or "
                    "configured regex pattern remains in extracted text."
                )
                st.warning(
                    "This is not a complete security certification. "
                    "Images, metadata, comments, tracked changes, "
                    "embedded content and missed detections may remain."
                )
                st.download_button(
                    "Download redacted document",
                    data=download_data,
                    file_name=f"{Path(uploaded.name).stem}_redacted{suffix}",
                    mime=MIME_TYPES[suffix],
                )

        except Exception as exc:
            duration = time.perf_counter() - started
            verification_status = "NOT COMPLETED"
            processing_status = "FAILED"

            try:
                save_history(
                    uploaded.name,
                    suffix.upper().lstrip(".") or "UNKNOWN",
                    file_size,
                    len(candidates),
                    operations,
                    categories,
                    duration,
                    processing_status,
                    verification_status,
                )
            except Exception as log_error:
                st.error(f"Could not save history: {log_error}")

            st.error(f"Processing or verification failed: {exc}")
            st.warning(
                "Download is unavailable because verification did not finish."
            )
            st.write(f"Time before failure: {duration:.2f} seconds")


# -------------------- PERSISTENT PROCESSING HISTORY --------------------

st.divider()
st.subheader("📚 Processing History")
st.caption(
    "History is stored locally in redaction_history.db. "
    "Document contents and detected sensitive values are not stored."
)

history = load_history()

if history:
    import pandas as pd

    columns = [
        "Timestamp",
        "Filename",
        "Type",
        "Size (bytes)",
        "Candidates",
        "Operations",
        "Input categories",
        "Duration (sec)",
        "Processing",
        "Verification",
    ]
    history_df = pd.DataFrame(history, columns=columns)
    st.dataframe(history_df, use_container_width=True, hide_index=True)

    st.download_button(
        "Export history as CSV",
        data=history_df.to_csv(index=False).encode("utf-8"),
        file_name="redaction_history.csv",
        mime="text/csv",
    )

    if st.button("Clear processing history"):
        with sqlite3.connect(DB_PATH) as conn:
            conn.execute("DELETE FROM redaction_history")
        st.success("Processing history cleared.")
        st.rerun()
else:
    st.write("No processing history yet. Redact a test document to begin.")
